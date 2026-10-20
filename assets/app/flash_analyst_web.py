"""Flash Analyst engine for the browser version, run with Pyodide.

GENERATED FILE: edit the flash_analyst/ package, then run tools/build_single_file.py.
"""

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import csv
import json
import re
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

__version__ = "1.0.0"

# ======================================================================
# DATA
# ======================================================================
# Canonical field name -> human label, in display order.
FIELDS = {
    "year": "Year",
    "revenue": "Revenue",
    "cogs": "Cost of Goods Sold",
    "operating_expenses": "Operating Expenses",
    "interest_expense": "Interest Expense",
    "taxes": "Taxes",
    "net_income": "Net Income",
    "total_assets": "Total Assets",
    "total_liabilities": "Total Liabilities",
    "shareholder_equity": "Shareholder Equity",
    "current_assets": "Current Assets",
    "current_liabilities": "Current Liabilities",
    "cash": "Cash",
    "operating_cash_flow": "Operating Cash Flow",
    "capex": "Capital Expenditures",
}

REQUIRED = ("year", "revenue", "net_income")

# Common alternative spellings people use in their own spreadsheets.
ALIASES = {
    "fiscal_year": "year",
    "period": "year",
    "sales": "revenue",
    "total_revenue": "revenue",
    "cost_of_goods_sold": "cogs",
    "cost_of_sales": "cogs",
    "cost_of_revenue": "cogs",
    "opex": "operating_expenses",
    "sg&a": "operating_expenses",
    "interest": "interest_expense",
    "tax": "taxes",
    "income_tax": "taxes",
    "net_profit": "net_income",
    "profit": "net_income",
    "assets": "total_assets",
    "liabilities": "total_liabilities",
    "equity": "shareholder_equity",
    "shareholders_equity": "shareholder_equity",
    "stockholders_equity": "shareholder_equity",
    "cash_and_equivalents": "cash",
    "cfo": "operating_cash_flow",
    "cash_from_operations": "operating_cash_flow",
    "capital_expenditures": "capex",
}

# Field label -> canonical key, so "Cost of Goods Sold" etc. are recognized too.

# Where to find each number, shown in the Excel template.
WHERE = {
    "revenue": "Income statement — top line (also called Sales)",
    "cogs": "Income statement — direct cost of the products/services sold",
    "operating_expenses": "Income statement — salaries, rent, marketing, admin (SG&A)",
    "interest_expense": "Income statement — interest paid on loans",
    "taxes": "Income statement — income tax expense",
    "net_income": "Income statement — bottom line (profit after everything)",
    "total_assets": "Balance sheet — total assets",
    "total_liabilities": "Balance sheet — total liabilities (everything owed)",
    "shareholder_equity": "Balance sheet — owners' / shareholders' equity",
    "current_assets": "Balance sheet — cash, receivables, inventory (due within 1 year)",
    "current_liabilities": "Balance sheet — bills and debts due within 1 year",
    "cash": "Balance sheet — cash and cash equivalents",
    "operating_cash_flow": "Cash flow statement — cash from operating activities",
    "capex": "Cash flow statement — purchases of equipment/property (enter as positive)",
}

# Total/summary lines as they appear in reports exported from accounting apps
# (QuickBooks, Xero, Wave, FreshBooks, Sage, Zoho Books, ...). These win over
# single-account lines with similar names, e.g. "Total Income" beats an account
# called "Sales".
EXPORT_LABELS = {
    "revenue": [
        "total income", "total revenue", "total revenues", "total sales", "total trading income",
        "total operating revenue", "net sales", "net revenue", "turnover", "total turnover",
    ],
    "cogs": [
        "total cost of goods sold", "total cost of sales", "total cost of revenue",
        "total direct costs", "total cogs", "cost of goods sold", "cost of sales",
    ],
    "operating_expenses": [
        "total expenses", "total operating expenses", "total overheads",
        "total selling general and administrative expenses", "total sg a",
    ],
    "interest_expense": ["interest expense", "total interest expense", "interest paid", "finance costs"],
    "taxes": ["income tax expense", "income taxes", "provision for income taxes", "corporation tax"],
    "net_income": [
        "net income", "net profit", "net profit loss", "net income loss", "net earnings",
        "profit for the year", "profit for the period", "net profit after tax",
    ],
    "total_assets": ["total assets"],
    "total_liabilities": ["total liabilities"],
    "shareholder_equity": [
        "total equity", "total shareholders equity", "total stockholders equity",
        "total owners equity", "total owner s equity", "net assets",
    ],
    "current_assets": ["total current assets"],
    "current_liabilities": ["total current liabilities"],
    "cash": [
        "total bank accounts", "total bank", "total cash", "total cash and cash equivalents",
        "cash and cash equivalents", "cash at bank and on hand",
    ],
    "operating_cash_flow": [
        "net cash provided by operating activities", "net cash from operating activities",
        "net cash flows from operating activities", "net cash used in operating activities",
        "cash from operating activities", "net cash provided by used in operating activities",
    ],
    "capex": [
        "purchase of property and equipment", "purchases of property and equipment",
        "purchase of property plant and equipment", "purchases of property plant and equipment",
        "purchase of fixed assets", "capital expenditures",
    ],
}

REPORT_TITLES = ("profit and loss", "profit & loss", "income statement", "balance sheet",
                 "cash flow", "statement of", "p&l")


def _label_key(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


_EXPORT_LOOKUP = {_label_key(label): key for key, labels in EXPORT_LABELS.items() for label in labels}


class DataError(ValueError):
    """Raised when an input file cannot be turned into a usable dataset."""


@dataclass
class CompanyData:
    name: str
    # One dict per year, sorted oldest -> newest. Missing values are None.
    rows: list = field(default_factory=list)

    @property
    def years(self):
        return [r["year"] for r in self.rows]

    def series(self, key):
        return [r.get(key) for r in self.rows]

    def has(self, key):
        return all(r.get(key) is not None for r in self.rows)

    @property
    def available_fields(self):
        return [k for k in FIELDS if any(r.get(k) is not None for r in self.rows)]


def normalize_header(header):
    key = str(header).strip().rstrip("*:").strip().lower().replace(" ", "_").replace("-", "_").replace("'", "")
    return ALIASES.get(key, key)


def parse_number(raw):
    """Parse '1,234', '$1.2M', '(500)', '12%' style values. Returns None for blanks."""
    if raw is None:
        return None
    text = str(raw).strip()
    if text in ("", "-", "—", "n/a", "N/A", "NA"):
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace("$", "").replace(",", "").replace("€", "").replace("£", "")
    multiplier = 1
    suffixes = {"k": 1e3, "m": 1e6, "b": 1e9, "t": 1e12}
    if text and text[-1].lower() in suffixes:
        multiplier = suffixes[text[-1].lower()]
        text = text[:-1]
    if text.endswith("%"):
        multiplier = 0.01
        text = text[:-1]
    try:
        value = float(text) * multiplier
    except ValueError:
        raise DataError(f"Could not read number: {raw!r}")
    return -value if negative else value

def _canonical(label):
    """Map a header/label cell to a FIELDS key, or None."""
    return _match(label)[0]


def _match(label):
    """Return (field key, priority). Report total lines rank above plain names."""
    if label is None or str(label).strip() == "":
        return None, 0
    export_key = _EXPORT_LOOKUP.get(_label_key(label))
    if export_key:
        return export_key, 2
    key = normalize_header(label)
    if key in FIELDS:
        return key, 1
    for field_key, field_label in FIELDS.items():
        if normalize_header(field_label) == key:
            return field_key, 1
    return None, 0


def _as_year(value):
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip().upper()
    for prefix in ("FY", "FY ", "CY"):
        if text.startswith(prefix):
            text = text[len(prefix):].strip()
    try:
        number = float(text)
    except ValueError:
        return None
    if number.is_integer() and 1900 <= number <= 2200:
        return int(number)
    return None


def _header_year(value):
    """Year from a column header: 2024, 'FY2024', 'Jan - Dec 2024', '31 Dec 2024', 'Dec-24'."""
    if isinstance(value, (datetime, date)):
        return value.year
    year = _as_year(value)
    if year or not isinstance(value, str):
        return year
    text = value.strip()
    if re.search(r"%|change|var|budget|forecast", text, re.I):
        return None
    years = re.findall(r"(?<!\d)((?:19|20)\d{2})(?!\d)", text)
    if years:
        return int(years[-1])
    short = re.fullmatch(r"[A-Za-z]{3,9}[-\s']+(\d{2})", text)
    return 2000 + int(short.group(1)) if short else None


def _cell_number(value, where):
    if isinstance(value, bool):
        raise DataError(f"{where}: expected a number")
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return parse_number(value)
    except DataError as exc:
        raise DataError(f"{where}: {exc}")


def _is_blank(value):
    return value is None or str(value).strip() == ""


def _find_name(grid):
    for row in grid:
        cells = [c for c in row if not _is_blank(c)]
        if len(cells) >= 2 and normalize_header(cells[0]) in ("company", "company_name", "name"):
            return str(cells[1]).strip()
    # Accounting exports: the company name is a lone line just above (QuickBooks) or
    # just below (Xero) the report title.
    def lone_text(row):
        cells = [c for c in row if not _is_blank(c)]
        return cells[0].strip() if len(cells) == 1 and isinstance(cells[0], str) else None

    def is_title(text):
        return bool(text) and any(t in text.lower() for t in REPORT_TITLES)

    def is_period(text):
        return bool(re.search(r"(19|20)\d{2}|^as (of|at)\b|^for the\b", text, re.I))

    lines = [lone_text(row) for row in grid[:6]]
    for i, text in enumerate(lines):
        if not is_title(text):
            continue
        for j in (i - 1, i + 1):
            if 0 <= j < len(lines) and lines[j] and not is_title(lines[j]) and not is_period(lines[j]):
                return lines[j]
    return None


def _records_years_down(grid, source):
    """Layout A: one header row (year, revenue, ...) then one row per year."""
    for h, header in enumerate(grid[:15]):
        mapping = {i: _canonical(c) for i, c in enumerate(header)}
        mapping = {i: k for i, k in mapping.items() if k}
        if set(REQUIRED) <= set(mapping.values()):
            rows = []
            for offset, raw in enumerate(grid[h + 1:], start=h + 2):
                if all(_is_blank(c) for c in raw):
                    continue
                row = {}
                for i, key in mapping.items():
                    value = raw[i] if i < len(raw) else None
                    row[key] = _cell_number(value, f"{source} row {offset}, column '{header[i]}'")
                if row.get("year") is None:
                    raise DataError(f"{source} row {offset}: year is blank")
                rows.append(row)
            return rows
    return None


def _records_years_across(grid, source):
    """Layout B: a header row of years across the top, one line item per row.

    Covers the template and most accounting-app exports. Returns whatever line items it
    recognizes (possibly without revenue, e.g. a balance sheet); load_company checks the rest.
    """
    for h, header in enumerate(grid[:20]):
        year_cols = {i: _header_year(c) for i, c in enumerate(header) if i > 0}
        year_cols = {i: y for i, y in year_cols.items() if y is not None}
        if not year_cols:
            # Single-period export: a "Total" column, with the period in a title line above.
            totals = [i for i, c in enumerate(header) if i > 0 and _label_key(c or "") == "total"]
            title_years = [_header_year(c) for row in grid[:h] for c in row if isinstance(c, str)]
            title_years = [y for y in title_years if y]
            if len(totals) == 1 and title_years:
                year_cols = {totals[0]: title_years[-1]}
            else:
                continue
        if len(set(year_cols.values())) != len(year_cols):
            dup = next(y for y in year_cols.values() if list(year_cols.values()).count(y) > 1)
            raise DataError(
                f"{source} has more than one column for {dup} (monthly or quarterly?). "
                f"Export the report with one column per year, e.g. 'Jan–Dec' totals."
            )
        found, priority, line_of = {}, {}, {}
        current_assets_heading = None
        first_col = min(year_cols)
        for offset, raw in enumerate(grid[h + 1:], start=h + 2):
            label = next((c for c in raw[:first_col] if not _is_blank(c)), None)
            if label is not None and _label_key(label) == "current assets" and current_assets_heading is None:
                current_assets_heading = offset
            key, rank = _match(label)
            if not key or key == "year" or rank <= priority.get(key, 0):
                continue
            values = {
                y: _cell_number(raw[i] if i < len(raw) else None, f"{source} row {offset} ({label}), {y}")
                for i, y in year_cols.items()
            }
            if all(v is None for v in values.values()):
                continue  # a section heading like "Assets" with no numbers
            found[key], priority[key], line_of[key] = values, rank, offset
        if not found:
            continue
        # Xero-style balance sheets list "Bank" as its own section *before* "Current Assets",
        # so Total Current Assets leaves the bank balance out. QuickBooks nests it inside.
        if ("cash" in found and "current_assets" in found and current_assets_heading
                and line_of["cash"] < current_assets_heading):
            found["current_assets"] = {
                y: None if v is None else v + (found["cash"][y] or 0)
                for y, v in found["current_assets"].items()
            }
        rows = []
        for y in year_cols.values():
            row = {"year": y}
            row.update({k: values[y] for k, values in found.items()})
            if all(row.get(k) is None for k in found):
                continue  # an empty year column the user didn't need
            rows.append(row)
        return rows
    return None


def _read_grid(path):
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, data_only=True, read_only=True)
        # Prefer a sheet called "Financials", then try the rest in order.
        sheets = sorted(wb.worksheets, key=lambda ws: ws.title.lower() != "financials")
        return [(f"{path.name} [{ws.title}]", [list(r) for r in ws.iter_rows(values_only=True)])
                for ws in sheets]
    if path.suffix.lower() in (".csv", ".txt"):
        with path.open(newline="", encoding="utf-8-sig") as fh:
            return [(path.name, [row for row in csv.reader(fh)])]
    raise DataError(f"{path.name}: unsupported file type (use .xlsx or .csv)")


def _load_partial(path):
    """Read every sheet of one file. Returns ({year: row}, company name or None)."""
    path = Path(path)
    if not path.exists():
        raise DataError(f"File not found: {path}")
    by_year, found_name = {}, None
    for source, grid in _read_grid(path):
        grid = [list(r) for r in grid if r is not None]
        rows = _records_years_down(grid, source) or _records_years_across(grid, source)
        if not rows:
            continue
        years = [int(r["year"]) for r in rows]
        if len(set(years)) != len(years):
            raise DataError(f"{source} has duplicate years")
        found_name = found_name or _find_name(grid)
        for row in rows:
            year = int(row["year"])
            merged = by_year.setdefault(year, {"year": year})
            for k, v in row.items():
                if v is not None and merged.get(k) is None:
                    merged[k] = v
    return by_year, found_name


def safe_text(text, limit=100):
    """Text from a customer's file, made safe to put in a cell: no control characters, and no leading
    = + - @ that Excel would run as a formula (formula injection)."""
    text = "".join(ch if ch.isprintable() else " " for ch in str(text))
    text = " ".join(text.split())[:limit]
    return text.lstrip("=+-@").strip()


def load_company(paths, name=None, default_name=None):
    """Combine one or more files (e.g. a P&L export and a balance sheet export) into one company."""
    paths = [Path(p) for p in paths]
    if not paths:
        raise DataError("No files given")
    by_year, found_name = {}, None
    for path in paths:
        part, part_name = _load_partial(path)
        found_name = found_name or part_name
        for year, row in part.items():
            merged = by_year.setdefault(year, {"year": year})
            for k, v in row.items():
                if v is not None and merged.get(k) is None:
                    merged[k] = v
    label = paths[0].name if len(paths) == 1 else f"{paths[0].parent.name or 'files'}/"
    if not by_year:
        raise DataError(
            f"{label}: couldn't find any numbers. It needs years plus Revenue and Net Income, either "
            f"as columns or as rows (an Excel/CSV export of a Profit & Loss report works). "
            f"Run `flash-analyst template` for a ready-to-fill file."
        )
    rows = [by_year[y] for y in sorted(by_year)]
    for r in rows:
        if r.get("capex") is not None:
            r["capex"] = abs(r["capex"])  # cash flow statements show it as negative
    # Years that only appear in a balance sheet (no P&L) can't be analyzed; drop them
    # if other years are complete, otherwise explain what's missing.
    complete = [r for r in rows if all(r.get(k) is not None for k in REQUIRED)]
    if not complete:
        missing = [FIELDS[k] for k in REQUIRED if all(r.get(k) is None for r in rows)]
        hint = (" Add the Profit & Loss (income statement) export too."
                if missing and len(missing) == len(REQUIRED) - 1 else "")
        r = next(r for r in rows if any(r.get(k) is None for k in REQUIRED))
        miss = ", ".join(FIELDS[k] for k in REQUIRED if r.get(k) is None)
        raise DataError(f"{label}: year {r['year']} is missing {miss}.{hint}")
    partial = [r for r in rows if r not in complete]
    if partial and any(r.get("revenue") is not None or r.get("net_income") is not None for r in partial):
        r = partial[0]
        miss = ", ".join(FIELDS[k] for k in REQUIRED if r.get(k) is None)
        raise DataError(f"{label}: year {r['year']} is missing {miss}")
    fallback = paths[0].stem.replace("_", " ").replace("-", " ").title()
    company = safe_text(name or found_name or default_name or fallback) or safe_text(fallback) or "Company"
    return CompanyData(name=company, rows=complete)


def load_file(path, name=None):
    """Load a company from one .csv or .xlsx file."""
    return load_company([path], name)


# Kept for backwards compatibility.
load_csv = load_file


def write_template(path, years=None):
    """Write a ready-to-fill template. .xlsx (default) is friendly; .csv is minimal."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            writer.writerow(list(FIELDS))
            writer.writerow([2024, 1000000, 400000, 350000, 20000, 50000, 180000,
                             1500000, 600000, 900000, 500000, 250000, 200000, 220000, 60000])
        return path
    return _write_xlsx_template(path, years)


def _write_xlsx_template(path, years=None):
    if years is None:
        last = date.today().year - 1
        years = list(range(last - 4, last + 1))
    navy = "1F3A5F"
    header_fill = PatternFill("solid", fgColor=navy)
    input_fill = PatternFill("solid", fgColor="FFF8E1")
    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    white_bold = Font(bold=True, color="FFFFFF")

    wb = Workbook()
    ws = wb.active
    ws.title = "Financials"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Company name:"
    ws["A1"].font = Font(bold=True)
    ws["B1"].fill = input_fill
    ws["B1"].border = border
    ws["B1"].font = Font(bold=True, size=12, color="1565C0")
    ws["A2"] = ("Fill in the yellow cells. Required: Revenue and Net Income. "
                "Everything else is optional but makes the report richer. Leave unused years blank.")
    ws["A2"].font = Font(italic=True, size=10, color="666666")

    header = ["Line item", "Where to find it"] + years
    for c, v in enumerate(header, start=1):
        cell = ws.cell(row=4, column=c, value=v)
        cell.fill, cell.font, cell.border = header_fill, white_bold, border
        cell.alignment = Alignment(horizontal="center")
    ws.cell(row=4, column=3).comment = Comment("You can change these years.", "Flash Analyst")

    r = 5
    for key, label in FIELDS.items():
        if key == "year":
            continue
        required = key in REQUIRED
        a = ws.cell(row=r, column=1, value=label + (" *" if required else ""))
        a.font = Font(bold=required)
        a.border = border
        b = ws.cell(row=r, column=2, value=WHERE.get(key, ""))
        b.font = Font(size=9, color="666666")
        b.border = border
        for c in range(3, 3 + len(years)):
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.border = input_fill, border
            cell.number_format = '#,##0;(#,##0)'
        r += 1
    ws.cell(row=r + 1, column=1, value="* required").font = Font(italic=True, size=9, color="666666")
    ws.cell(row=r + 2, column=1, value=("Tips: enter whole numbers (e.g. 1250000), use a minus sign for "
                                        "losses, and enter Capital Expenditures as a positive number."))
    ws.cell(row=r + 2, column=1).font = Font(italic=True, size=9, color="666666")

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 58
    for c in range(3, 3 + len(years)):
        ws.column_dimensions[ws.cell(row=4, column=c).column_letter].width = 15
    ws.freeze_panes = "C5"
    wb.save(path)
    return path


# ======================================================================
# ANALYSIS
# ======================================================================
# Ratio key -> (label, number format). "pct" = percentage, "x" = multiple, "money" = currency.
RATIOS = {
    "gross_margin": ("Gross Margin", "pct"),
    "operating_margin": ("Operating Margin", "pct"),
    "net_margin": ("Net Margin", "pct"),
    "revenue_growth": ("Revenue Growth (YoY)", "pct"),
    "net_income_growth": ("Net Income Growth (YoY)", "pct"),
    "roa": ("Return on Assets", "pct"),
    "roe": ("Return on Equity", "pct"),
    "current_ratio": ("Current Ratio", "x"),
    "debt_to_equity": ("Debt to Equity", "x"),
    "interest_coverage": ("Interest Coverage", "x"),
    "free_cash_flow": ("Free Cash Flow", "money"),
    "fcf_margin": ("FCF Margin", "pct"),
}

# Score component -> (label, weight).
SCORE_WEIGHTS = {
    "profitability": ("Profitability", 25),
    "growth": ("Growth", 25),
    "liquidity": ("Liquidity", 20),
    "leverage": ("Leverage", 15),
    "cash": ("Cash Generation", 15),
}

# Lowest score for each grade; anything below the D cutoff is an F.
DEFAULT_GRADES = (("A", 80), ("B", 65), ("C", 50), ("D", 35))


@dataclass
class Analysis:
    company: object
    ratios: list = field(default_factory=list)  # one dict per year, aligned with company.rows
    revenue_cagr: float = None
    score: float = None
    grade: str = None
    grade_cutoffs: tuple = DEFAULT_GRADES
    components: dict = field(default_factory=dict)  # key -> 0..1, or None if not scorable
    insights: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    forecast: list = field(default_factory=list)


def _div(a, b):
    if a is None or b is None or b == 0:
        return None
    return a / b


def _growth(curr, prev):
    if curr is None or prev is None or prev == 0:
        return None
    return (curr - prev) / abs(prev)


def _clamp(x, lo=0.0, hi=1.0):
    return max(lo, min(hi, x))


def cagr(first, last, periods):
    if first is None or last is None or periods <= 0 or first <= 0 or last <= 0:
        return None
    return (last / first) ** (1 / periods) - 1


def year_ratios(row, prev=None):
    rev = row.get("revenue")
    cogs = row.get("cogs")
    opex = row.get("operating_expenses")
    gross = rev - cogs if rev is not None and cogs is not None else None
    op_income = gross - opex if gross is not None and opex is not None else None
    fcf = (row["operating_cash_flow"] - row["capex"]
           if row.get("operating_cash_flow") is not None and row.get("capex") is not None else None)
    return {
        "gross_margin": _div(gross, rev),
        "operating_margin": _div(op_income, rev),
        "net_margin": _div(row.get("net_income"), rev),
        "revenue_growth": _growth(rev, prev.get("revenue")) if prev else None,
        "net_income_growth": _growth(row.get("net_income"), prev.get("net_income")) if prev else None,
        "roa": _div(row.get("net_income"), row.get("total_assets")),
        "roe": _div(row.get("net_income"), row.get("shareholder_equity")),
        "current_ratio": _div(row.get("current_assets"), row.get("current_liabilities")),
        "debt_to_equity": _div(row.get("total_liabilities"), row.get("shareholder_equity")),
        "interest_coverage": _div(op_income, row.get("interest_expense")),
        "free_cash_flow": fcf,
        "fcf_margin": _div(fcf, rev),
    }


def score_components(latest, revenue_cagr):
    """Each component is scaled 0..1; None means the data needed is missing."""
    c = {}
    nm = latest.get("net_margin")
    c["profitability"] = None if nm is None else _clamp(nm / 0.20)
    c["growth"] = None if revenue_cagr is None else _clamp(revenue_cagr / 0.15)
    cr = latest.get("current_ratio")
    c["liquidity"] = None if cr is None else _clamp(cr / 2.0)
    de = latest.get("debt_to_equity")
    if de is None:
        c["leverage"] = None
    elif de < 0:  # negative equity
        c["leverage"] = 0.0
    else:
        c["leverage"] = _clamp(1 - (de - 0.5) / 2.5)
    fm = latest.get("fcf_margin")
    c["cash"] = None if fm is None else _clamp(fm / 0.15)
    return c


def overall_score(components):
    total = weight = 0
    for key, (_, w) in SCORE_WEIGHTS.items():
        if components.get(key) is not None:
            total += components[key] * w
            weight += w
    if weight == 0:
        return None
    return round(100 * total / weight, 1)


def parse_grades(text):
    """Turn "A=85,B=70,C=55,D=40" into grade cutoffs. Letters left out keep their usual cutoff."""
    cutoffs = dict(DEFAULT_GRADES)
    for part in filter(None, (p.strip() for p in text.replace(";", ",").split(","))):
        letter, sep, value = part.partition("=")
        letter = letter.strip().upper()
        if not sep or letter not in cutoffs:
            raise ValueError(f"'{part}' isn't a grade setting. Use A, B, C and D, like A=85,B=70,C=55,D=40")
        try:
            cutoffs[letter] = float(value)
        except ValueError:
            raise ValueError(f"'{value.strip()}' isn't a number (in '{part}')")
        if not 0 < cutoffs[letter] <= 100:
            raise ValueError(f"Grade {letter} must be between 0 and 100")
    values = [cutoffs[letter] for letter, _ in DEFAULT_GRADES]
    if any(a <= b for a, b in zip(values, values[1:])):
        shown = ", ".join(f"{letter}={_num(cutoffs[letter])}" for letter, _ in DEFAULT_GRADES)
        raise ValueError(f"Each grade must need a higher score than the next one (A > B > C > D); got {shown}")
    return tuple((letter, cutoffs[letter]) for letter, _ in DEFAULT_GRADES)


def _num(x):
    return f"{x:g}"


def grade_scale_text(cutoffs=DEFAULT_GRADES):
    """E.g. "A 80+ · B 65+ · C 50+ · D 35+ · F below 35"."""
    parts = [f"{letter} {_num(cutoff)}+" for letter, cutoff in cutoffs]
    return " · ".join(parts + [f"F below {_num(cutoffs[-1][1])}"])


def grade_for(score, cutoffs=DEFAULT_GRADES):
    if score is None:
        return "N/A"
    for letter, cutoff in cutoffs:
        if score >= cutoff:
            return letter
    return "F"


def fmt_pct(x):
    return f"{x * 100:.1f}%"


def fmt_money(x):
    sign = "-" if x < 0 else ""
    x = abs(x)
    for div, suffix in ((1e12, "T"), (1e9, "B"), (1e6, "M"), (1e3, "K")):
        if x >= div:
            return f"{sign}${x / div:.1f}{suffix}"
    return f"{sign}${x:,.0f}"


def build_insights(company, ratios, revenue_cagr):
    insights, warnings = [], []
    rows = company.rows
    first, last = rows[0], rows[-1]
    latest = ratios[-1]
    span = last["year"] - first["year"]

    if revenue_cagr is not None and span > 0:
        trend = "grew" if revenue_cagr >= 0 else "shrank"
        insights.append(
            f"Revenue {trend} from {fmt_money(first['revenue'])} to {fmt_money(last['revenue'])} "
            f"({fmt_pct(revenue_cagr)} CAGR over {span} year{'s' if span != 1 else ''})."
        )
    g = latest.get("revenue_growth")
    if g is not None and g < 0:
        warnings.append(f"Revenue fell {fmt_pct(-g)} in {last['year']}.")

    nm_first, nm_last = ratios[0].get("net_margin"), latest.get("net_margin")
    if nm_last is not None:
        if nm_first is not None and len(rows) > 1:
            delta = nm_last - nm_first
            direction = "expanded" if delta > 0.005 else "contracted" if delta < -0.005 else "held steady"
            insights.append(
                f"Net margin {direction} from {fmt_pct(nm_first)} to {fmt_pct(nm_last)}."
            )
        else:
            insights.append(f"Net margin is {fmt_pct(nm_last)}.")
        if nm_last < 0:
            warnings.append(f"The company lost money in {last['year']} (net margin {fmt_pct(nm_last)}).")

    gm = latest.get("gross_margin")
    if gm is not None:
        insights.append(f"Gross margin of {fmt_pct(gm)} in {last['year']}.")

    roe = latest.get("roe")
    if roe is not None:
        quality = "strong" if roe >= 0.15 else "moderate" if roe >= 0.08 else "weak"
        insights.append(f"Return on equity is {quality} at {fmt_pct(roe)}.")

    cr = latest.get("current_ratio")
    if cr is not None:
        if cr < 1:
            warnings.append(f"Current ratio of {cr:.2f}x — short-term liabilities exceed current assets.")
        else:
            insights.append(f"Current ratio of {cr:.2f}x covers short-term obligations.")

    de = latest.get("debt_to_equity")
    if de is not None:
        if de < 0:
            warnings.append("Shareholder equity is negative.")
        elif de > 2:
            warnings.append(f"High leverage: debt-to-equity of {de:.2f}x.")
        else:
            insights.append(f"Debt-to-equity of {de:.2f}x.")

    ic = latest.get("interest_coverage")
    if ic is not None and ic < 3:
        warnings.append(f"Interest coverage is thin at {ic:.1f}x operating income.")

    fcf = latest.get("free_cash_flow")
    if fcf is not None:
        if fcf < 0:
            warnings.append(f"Negative free cash flow of {fmt_money(fcf)} in {last['year']}.")
        else:
            insights.append(f"Generated {fmt_money(fcf)} of free cash flow in {last['year']}.")

    return insights, warnings


def build_forecast(company, ratios, revenue_cagr, years=3):
    """Simple projection: revenue at historical CAGR, net income at average recent margin."""
    last = company.rows[-1]
    growth = revenue_cagr
    if growth is None:
        growth = ratios[-1].get("revenue_growth") or 0.0
    growth = max(-0.5, min(1.0, growth))
    margins = [r["net_margin"] for r in ratios[-3:] if r.get("net_margin") is not None]
    margin = sum(margins) / len(margins) if margins else 0.0
    out, revenue = [], last["revenue"]
    for i in range(1, years + 1):
        revenue *= 1 + growth
        out.append({"year": last["year"] + i, "revenue": revenue, "net_income": revenue * margin})
    return growth, margin, out


def analyze(company, forecast_years=3, grade_cutoffs=DEFAULT_GRADES):
    ratios, prev = [], None
    for row in company.rows:
        ratios.append(year_ratios(row, prev))
        prev = row
    first, last = company.rows[0], company.rows[-1]
    rc = cagr(first["revenue"], last["revenue"], last["year"] - first["year"])
    components = score_components(ratios[-1], rc)
    score = overall_score(components)
    insights, warnings = build_insights(company, ratios, rc)
    _, _, forecast = build_forecast(company, ratios, rc, forecast_years)
    return Analysis(
        company=company,
        ratios=ratios,
        revenue_cagr=rc,
        score=score,
        grade=grade_for(score, grade_cutoffs),
        grade_cutoffs=grade_cutoffs,
        components=components,
        insights=insights,
        warnings=warnings,
        forecast=forecast,
    )


# ======================================================================
# NEEDS
# ======================================================================
PRIORITIES = ("Urgent", "High", "Medium", "Low")


@dataclass
class Need:
    priority: str      # one of PRIORITIES
    area: str          # e.g. "Cash & liquidity"
    finding: str       # what the numbers show, with the number
    need: str          # what the client needs to do
    service: str       # what an advisor could offer them
    target: str        # a measurable goal


def _rank(need):
    return PRIORITIES.index(need.priority)


def client_needs(analysis):
    """Rules of thumb that map weak (or strong) numbers to recommended actions."""
    company = analysis.company
    latest_row, latest = company.rows[-1], analysis.ratios[-1]
    year = latest_row["year"]
    needs = []

    def add(*args):
        needs.append(Need(*args))

    # Profitability
    nm = latest.get("net_margin")
    if nm is not None:
        if nm < 0:
            add("Urgent", "Profitability",
                f"Lost {fmt_money(-latest_row['net_income'])} in {year} (net margin {fmt_pct(nm)}).",
                "Get back to profit: find the biggest costs to cut and check prices cover costs.",
                "Profitability review and cost-reduction plan", "Net margin above 0%, then 5%+")
        elif nm < 0.05:
            add("High", "Profitability", f"Thin net margin of {fmt_pct(nm)} in {year}.",
                "Widen margins so a bad month doesn't turn into a loss.",
                "Pricing and cost review", "Net margin of 5–10%")
        elif nm < 0.10:
            add("Medium", "Profitability", f"Net margin of {fmt_pct(nm)} has room to grow.",
                "Tune pricing and overheads to keep more of each sale.",
                "Margin improvement plan", "Net margin of 10%+")

    gm = latest.get("gross_margin")
    if gm is not None and gm < 0.30:
        add("Medium", "Pricing & direct costs", f"Gross margin is only {fmt_pct(gm)}.",
            "Raise prices or lower supplier and production costs.",
            "Pricing strategy and supplier review", "Gross margin of 30%+")

    # Growth
    growth = latest.get("revenue_growth")
    if growth is not None and growth < 0:
        add("High", "Sales", f"Revenue fell {fmt_pct(-growth)} in {year}.",
            "Find out why sales dropped and win them back: lost customers, pricing, marketing.",
            "Sales recovery plan", "Revenue back above last year's level")
    elif analysis.revenue_cagr is not None and analysis.revenue_cagr < 0.03:
        add("Medium", "Growth", f"Revenue is growing only {fmt_pct(analysis.revenue_cagr)} a year.",
            "Set out where new sales will come from: new customers, products or markets.",
            "Growth strategy and budget", "Revenue growth of 5–10% a year")

    # Cash & liquidity
    cr = latest.get("current_ratio")
    if cr is not None:
        if cr < 1:
            add("Urgent", "Cash & liquidity",
                f"Current ratio of {cr:.2f}x: bills due this year exceed short-term assets.",
                "Build a week-by-week cash forecast, collect money owed faster and agree longer payment terms.",
                "13-week cash flow forecast and working-capital plan", "Current ratio of 1.2x+")
        elif cr < 1.5:
            add("Medium", "Cash & liquidity", f"Current ratio of {cr:.2f}x leaves little cushion.",
                "Build a cash buffer to cover slow months.",
                "Cash flow forecasting", "Current ratio of 1.5x+")

    fcf = latest.get("free_cash_flow")
    if fcf is not None:
        if fcf < 0:
            add("High", "Cash generation", f"Free cash flow was {fmt_money(fcf)} in {year}.",
                "Plan equipment and investment spending against the cash the business actually makes.",
                "Capital spending and cash plan", "Positive free cash flow")
        elif latest.get("fcf_margin") is not None and latest["fcf_margin"] < 0.05:
            add("Medium", "Cash generation", f"Free cash flow is only {fmt_pct(latest['fcf_margin'])} of revenue.",
                "Turn more profit into cash: tighten collections and stock levels.",
                "Working-capital review", "FCF margin of 5%+")

    # Debt
    de = latest.get("debt_to_equity")
    if de is not None:
        if de < 0:
            add("Urgent", "Debt & solvency", "Owners' equity is negative: the business owes more than it owns.",
                "Restructure debt and rebuild equity before taking on anything new.",
                "Debt restructuring and recovery plan", "Positive equity")
        elif de > 3:
            add("Urgent", "Debt & solvency", f"Debt-to-equity of {de:.2f}x is very high.",
                "Reduce or refinance debt; avoid new borrowing.",
                "Debt restructuring / refinancing", "Debt-to-equity under 2x")
        elif de > 2:
            add("High", "Debt & solvency", f"Debt-to-equity of {de:.2f}x is high.",
                "Pay down debt and check loan terms.",
                "Refinancing review", "Debt-to-equity under 1.5x")

    ic = latest.get("interest_coverage")
    if ic is not None:
        if ic < 1.5:
            add("Urgent", "Debt & solvency",
                f"Operating profit covers interest only {ic:.1f}x.",
                "Cut interest costs: refinance, renegotiate or repay the most expensive loans.",
                "Lender negotiation and refinancing", "Interest coverage of 3x+")
        elif ic < 3:
            add("High", "Debt & solvency", f"Interest coverage of {ic:.1f}x is thin.",
                "Lower interest costs before rates or profits move against the business.",
                "Refinancing review", "Interest coverage of 3x+")

    roe = latest.get("roe")
    if roe is not None and 0 <= roe < 0.08:
        add("Medium", "Returns", f"Return on equity is {fmt_pct(roe)}.",
            "Make the owners' money work harder: drop low-return activities, use idle assets.",
            "Business performance review", "Return on equity of 10%+")

    # Missing records
    missing = []
    if not company.has("total_assets"):
        missing.append("balance sheet")
    if not company.has("operating_cash_flow"):
        missing.append("cash flow statement")
    if missing:
        add("Low", "Records", f"No {' or '.join(missing)} provided.",
            f"Provide the {' and '.join(missing)} for a complete health check.",
            "Bookkeeping clean-up and year-end accounts", "Full financial statements each year")

    # Opportunities for healthy businesses
    if not any(n.priority == "Urgent" for n in needs):
        if analysis.revenue_cagr is not None and analysis.revenue_cagr >= 0.10:
            add("Low", "Opportunity", f"Strong growth of {fmt_pct(analysis.revenue_cagr)} a year.",
                "Plan the next stage: hiring, systems and funding to keep growing smoothly.",
                "Scale-up plan and budget", "Keep growth above 10% without losing margin")
        if cr is not None and cr >= 2 and (de is None or de < 1):
            add("Low", "Opportunity", "Plenty of cash and little debt.",
                "Put spare cash to work: expansion, equipment or investments.",
                "Investment and expansion planning", "Agreed plan for surplus cash")
        if nm is not None and nm >= 0.10:
            add("Low", "Tax & owners", f"Healthy profit ({fmt_pct(nm)} net margin).",
                "Plan taxes and how profits are paid to owners or reinvested.",
                "Tax planning and owner pay review", "Tax plan agreed before year-end")

    needs.sort(key=_rank)  # stable: keeps the order above within each priority
    return needs


def overall_priority(needs):
    return needs[0].priority if needs else "Low"


# ======================================================================
# WORKBOOK
# ======================================================================
MONEY = '_($* #,##0_);_($* (#,##0);_($* "-"_);_(@_)'
PCT = "0.0%"
MULT = '0.00"x"'
FORMATS = {"pct": PCT, "x": MULT, "money": MONEY}

PRIORITY_COLORS = {"Urgent": "C62828", "High": "EF6C00", "Medium": "F9A825", "Low": "2E7D32"}
STATUSES = ("Not started", "In progress", "Done")

GRADE_COLORS = {"A": "2E7D32", "B": "558B2F", "C": "F9A825", "D": "EF6C00", "F": "C62828", "N/A": "757575"}


class Theme:
    def __init__(self, color="1F3A5F"):
        color = color.lstrip("#").upper()
        self.color = color
        self.header_fill = PatternFill("solid", fgColor=color)
        self.header_font = Font(bold=True, color="FFFFFF")
        self.title_font = Font(bold=True, size=18, color=color)
        self.subtitle_font = Font(italic=True, size=10, color="666666")
        self.section_font = Font(bold=True, size=12, color=color)
        self.band_fill = PatternFill("solid", fgColor="F3F6FA")
        self.input_fill = PatternFill("solid", fgColor="FFF8E1")
        thin = Side(style="thin", color="D0D7E2")
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _header_row(ws, row, values, theme, start_col=1):
    for i, v in enumerate(values):
        cell = ws.cell(row=row, column=start_col + i, value=v)
        cell.fill = theme.header_fill
        cell.font = theme.header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = theme.border


def _title(ws, text, subtitle, theme):
    ws["A1"] = text
    ws["A1"].font = theme.title_font
    ws["A2"] = subtitle
    ws["A2"].font = theme.subtitle_font
    ws.sheet_view.showGridLines = False


def _subtitle(analysis, prepared_by):
    parts = [f"Generated {date.today():%B %d, %Y}"]
    if prepared_by and safe_text(prepared_by):
        parts.append(f"Prepared by {safe_text(prepared_by)}")
    return " · ".join(parts)


def _financials_sheet(wb, analysis, theme, subtitle):
    """Years across columns, line items down rows. Returns {field: row} for formulas."""
    company = analysis.company
    ws = wb.create_sheet("Financials")
    _title(ws, f"{company.name} — Financial Statements", subtitle, theme)
    _header_row(ws, 4, ["Line Item"] + [str(y) for y in company.years], theme)
    rows = {}
    r = 5
    for key in company.available_fields:
        if key == "year":
            continue
        ws.cell(row=r, column=1, value=FIELDS[key]).border = theme.border
        for c, value in enumerate(company.series(key), start=2):
            cell = ws.cell(row=r, column=c, value=value)
            cell.number_format = MONEY
            cell.border = theme.border
            if r % 2 == 0:
                cell.fill = theme.band_fill
        if r % 2 == 0:
            ws.cell(row=r, column=1).fill = theme.band_fill
        rows[key] = r
        r += 1
    ws.column_dimensions["A"].width = 26
    for c in range(2, len(company.years) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.freeze_panes = "B5"
    return ws, rows


def _ratio_formula(key, col, prev_col, f):
    """Excel formula for a ratio; f(field, col) gives a Financials cell reference."""
    def has(*keys):
        return all(k in f.rows for k in keys)

    if key == "gross_margin" and has("revenue", "cogs"):
        return f"({f('revenue', col)}-{f('cogs', col)})/{f('revenue', col)}"
    if key == "operating_margin" and has("revenue", "cogs", "operating_expenses"):
        return (f"({f('revenue', col)}-{f('cogs', col)}-{f('operating_expenses', col)})"
                f"/{f('revenue', col)}")
    if key == "net_margin":
        return f"{f('net_income', col)}/{f('revenue', col)}"
    if key == "revenue_growth" and prev_col:
        return f"({f('revenue', col)}-{f('revenue', prev_col)})/ABS({f('revenue', prev_col)})"
    if key == "net_income_growth" and prev_col:
        return (f"({f('net_income', col)}-{f('net_income', prev_col)})"
                f"/ABS({f('net_income', prev_col)})")
    if key == "roa" and has("total_assets"):
        return f"{f('net_income', col)}/{f('total_assets', col)}"
    if key == "roe" and has("shareholder_equity"):
        return f"{f('net_income', col)}/{f('shareholder_equity', col)}"
    if key == "current_ratio" and has("current_assets", "current_liabilities"):
        return f"{f('current_assets', col)}/{f('current_liabilities', col)}"
    if key == "debt_to_equity" and has("total_liabilities", "shareholder_equity"):
        return f"{f('total_liabilities', col)}/{f('shareholder_equity', col)}"
    if key == "interest_coverage" and has("revenue", "cogs", "operating_expenses", "interest_expense"):
        return (f"({f('revenue', col)}-{f('cogs', col)}-{f('operating_expenses', col)})"
                f"/{f('interest_expense', col)}")
    if key == "free_cash_flow" and has("operating_cash_flow", "capex"):
        return f"{f('operating_cash_flow', col)}-{f('capex', col)}"
    if key == "fcf_margin" and has("operating_cash_flow", "capex"):
        return f"({f('operating_cash_flow', col)}-{f('capex', col)})/{f('revenue', col)}"
    return None


def _ratios_sheet(wb, analysis, fin_rows, theme, subtitle):
    company = analysis.company
    ws = wb.create_sheet("Ratios")
    _title(ws, f"{company.name} — Key Ratios", subtitle + " · Live formulas linked to Financials", theme)
    _header_row(ws, 4, ["Ratio"] + [str(y) for y in company.years], theme)

    def ref(field, col):
        return f"Financials!${col}${fin_rows[field]}"
    ref.rows = fin_rows

    ratio_rows = {}
    r = 5
    for key, (label, kind) in RATIOS.items():
        if all(yr.get(key) is None for yr in analysis.ratios):
            continue
        ws.cell(row=r, column=1, value=label).border = theme.border
        for i in range(len(company.years)):
            col = get_column_letter(i + 2)
            prev_col = get_column_letter(i + 1) if i > 0 else None
            formula = _ratio_formula(key, col, prev_col, ref)
            cell = ws.cell(row=r, column=i + 2)
            if formula:
                cell.value = f'=IFERROR({formula},"")'
            cell.number_format = FORMATS[kind]
            cell.border = theme.border
            if r % 2 == 0:
                cell.fill = theme.band_fill
        if r % 2 == 0:
            ws.cell(row=r, column=1).fill = theme.band_fill
        ratio_rows[key] = r
        r += 1
    ws.column_dimensions["A"].width = 26
    for c in range(2, len(company.years) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.freeze_panes = "B5"
    return ws, ratio_rows


def _forecast_sheet(wb, analysis, fin_rows, theme, subtitle):
    company = analysis.company
    growth, margin, projection = build_forecast(
        company, analysis.ratios, analysis.revenue_cagr, len(analysis.forecast) or 3)
    ws = wb.create_sheet("Forecast")
    _title(ws, f"{company.name} — Projection", subtitle, theme)

    ws["A4"] = "Assumptions (edit the yellow cells)"
    ws["A4"].font = theme.section_font
    ws["A5"], ws["B5"] = "Annual revenue growth", growth
    ws["A6"], ws["B6"] = "Net margin", margin
    for cell in ("B5", "B6"):
        ws[cell].number_format = PCT
        ws[cell].fill = theme.input_fill
        ws[cell].border = theme.border
        ws[cell].font = Font(bold=True, color="1565C0")

    _header_row(ws, 8, ["Year", "Revenue", "Net Income", "Type"], theme)
    r = 9
    for row in company.rows:
        ws.cell(row=r, column=1, value=row["year"])
        ws.cell(row=r, column=2, value=row["revenue"])
        ws.cell(row=r, column=3, value=row["net_income"])
        ws.cell(row=r, column=4, value="Actual")
        r += 1
    for _ in projection:
        ws.cell(row=r, column=1, value=f"=A{r - 1}+1")
        ws.cell(row=r, column=2, value=f"=B{r - 1}*(1+$B$5)")
        ws.cell(row=r, column=3, value=f"=B{r}*$B$6")
        ws.cell(row=r, column=4, value="Forecast")
        for c in range(1, 5):
            ws.cell(row=r, column=c).font = Font(italic=True, color="1565C0")
        r += 1
    for rr in range(9, r):
        ws.cell(row=rr, column=2).number_format = MONEY
        ws.cell(row=rr, column=3).number_format = MONEY
        for c in range(1, 5):
            ws.cell(row=rr, column=c).border = theme.border

    chart = LineChart()
    chart.title = "Revenue & Net Income: Actual + Forecast"
    chart.height, chart.width = 8, 18
    chart.y_axis.numFmt = '"$"#,##0'
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=8, max_row=r - 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=9, max_row=r - 1))
    ws.add_chart(chart, "F4")

    ws["A" + str(r + 1)] = ("Projection uses the historical revenue CAGR and the average net margin of the "
                            "last 3 years. It is an illustration, not investment advice.")
    ws["A" + str(r + 1)].font = theme.subtitle_font
    ws.column_dimensions["A"].width = 24
    for c in "BCD":
        ws.column_dimensions[c].width = 16
    return ws


def _performance_chart(fin_ws, fin_rows, n_years):
    chart = BarChart()
    chart.title = "Revenue vs Net Income"
    chart.height, chart.width = 8, 16
    chart.y_axis.numFmt = '"$"#,##0'
    for key in ("revenue", "net_income"):
        r = fin_rows[key]
        chart.add_data(Reference(fin_ws, min_col=1, max_col=n_years + 1, min_row=r, max_row=r),
                       from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(fin_ws, min_col=2, max_col=n_years + 1, min_row=4, max_row=4))
    return chart


def _margin_chart(ratio_ws, ratio_rows, n_years):
    keys = [k for k in ("gross_margin", "operating_margin", "net_margin", "fcf_margin") if k in ratio_rows]
    if not keys:
        return None
    chart = LineChart()
    chart.title = "Margin Trends"
    chart.height, chart.width = 8, 16
    chart.y_axis.numFmt = "0%"
    for key in keys:
        r = ratio_rows[key]
        chart.add_data(Reference(ratio_ws, min_col=1, max_col=n_years + 1, min_row=r, max_row=r),
                       from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(ratio_ws, min_col=2, max_col=n_years + 1, min_row=4, max_row=4))
    return chart


def _kpi(ws, row, col, label, value, theme):
    top = ws.cell(row=row, column=col, value=label)
    top.font = Font(size=9, bold=True, color="FFFFFF")
    top.fill = theme.header_fill
    top.alignment = Alignment(horizontal="center")
    bottom = ws.cell(row=row + 1, column=col, value=value)
    bottom.font = Font(size=16, bold=True, color=theme.color)
    bottom.alignment = Alignment(horizontal="center", vertical="center")
    bottom.fill = theme.band_fill
    ws.row_dimensions[row + 1].height = 30


def _dashboard_sheet(wb, analysis, theme, subtitle, charts):
    company = analysis.company
    latest_row, latest = company.rows[-1], analysis.ratios[-1]
    ws = wb.active
    ws.title = "Dashboard"
    _title(ws, f"{company.name} — Company Analysis", subtitle, theme)

    def show(value, fn):
        return "n/a" if value is None else fn(value)

    kpis = [
        (f"Revenue {latest_row['year']}", fmt_money(latest_row["revenue"])),
        ("Revenue CAGR", show(analysis.revenue_cagr, fmt_pct)),
        ("Net Margin", show(latest.get("net_margin"), fmt_pct)),
        ("ROE", show(latest.get("roe"), fmt_pct)),
        ("Current Ratio", show(latest.get("current_ratio"), lambda v: f"{v:.2f}x")),
        ("Debt / Equity", show(latest.get("debt_to_equity"), lambda v: f"{v:.2f}x")),
    ]
    for i, (label, value) in enumerate(kpis):
        _kpi(ws, 4, 1 + i, label, value, theme)

    ws["A7"] = "Financial Health Score"
    ws["A7"].font = theme.section_font
    ws["A8"] = "n/a" if analysis.score is None else analysis.score
    ws["A8"].font = Font(size=28, bold=True, color=GRADE_COLORS[analysis.grade])
    ws["B8"] = f"Grade {analysis.grade}"
    ws["B8"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["B8"].fill = PatternFill("solid", fgColor=GRADE_COLORS[analysis.grade])
    ws["B8"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[8].height = 40
    ws["C8"] = f"Grade scale: {grade_scale_text(analysis.grade_cutoffs)}"
    ws["C8"].font = Font(italic=True, color="757575")
    ws["C8"].alignment = Alignment(vertical="center")

    r = 10
    _header_row(ws, r, ["Component", "Weight", "Score"], theme)
    for key, (label, weight) in SCORE_WEIGHTS.items():
        r += 1
        value = analysis.components.get(key)
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=weight / 100).number_format = "0%"
        cell = ws.cell(row=r, column=3, value="no data" if value is None else value)
        cell.number_format = "0%"
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = theme.border
    ws.conditional_formatting.add(
        f"C11:C{r}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                 color=theme.color))

    r += 2
    ws.cell(row=r, column=1, value="Key Insights").font = theme.section_font
    for text in analysis.insights:
        r += 1
        ws.cell(row=r, column=1, value=f"✔  {text}")
    if analysis.warnings:
        r += 2
        ws.cell(row=r, column=1, value="Watch-outs").font = Font(bold=True, size=12, color="C62828")
        for text in analysis.warnings:
            r += 1
            ws.cell(row=r, column=1, value=f"⚠  {text}").font = Font(color="C62828")

    r += 2
    ws.cell(row=r, column=1, value=("Score = weighted blend of net margin, revenue CAGR, current ratio, "
                                    "debt-to-equity and FCF margin. For informational purposes only."))
    ws.cell(row=r, column=1).font = theme.subtitle_font

    for i in range(1, 7):
        ws.column_dimensions[get_column_letter(i)].width = 17
    anchor_row = 4
    for chart in charts:
        if chart is not None:
            ws.add_chart(chart, f"H{anchor_row}")
            anchor_row += 17
    return ws


def _client_needs_sheet(wb, analysis, theme, subtitle):
    """Pro: a prioritized action plan the advisor can work through with the client."""
    needs = client_needs(analysis)
    ws = wb.create_sheet("Client Needs", 1)
    _title(ws, f"{analysis.company.name} — What This Client Needs", subtitle, theme)

    counts = {p: sum(n.priority == p for n in needs) for p in PRIORITIES}
    overall = overall_priority(needs)
    ws["A4"] = "Overall priority"
    ws["A4"].font = theme.section_font
    ws["C4"] = overall
    ws["C4"].font = Font(bold=True, color="FFFFFF", size=12)
    ws["C4"].fill = PatternFill("solid", fgColor=PRIORITY_COLORS[overall])
    ws["C4"].alignment = Alignment(horizontal="center")
    ws["D4"] = " · ".join(f"{counts[p]} {p.lower()}" for p in PRIORITIES if counts[p]) or "No issues found"
    ws["A5"] = (f"Health score {analysis.score if analysis.score is not None else 'n/a'} (grade "
                f"{analysis.grade}). Fill in Owner, Due date and Status to track progress with the client.")
    ws["A5"].font = theme.subtitle_font

    headers = ["#", "Priority", "Area", "What we found", "What the client needs",
               "Suggested service", "Target", "Owner", "Due date", "Status"]
    _header_row(ws, 7, headers, theme)
    wrap = Alignment(wrap_text=True, vertical="top")
    r = 7
    for i, n in enumerate(needs, start=1):
        r += 1
        values = [i, n.priority, n.area, n.finding, n.need, n.service, n.target, None, None, STATUSES[0]]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = wrap
            cell.border = theme.border
            if c >= 8:
                cell.fill = theme.input_fill
        pcell = ws.cell(row=r, column=2)
        pcell.font = Font(bold=True, color="1A1A1A" if n.priority == "Medium" else "FFFFFF")
        pcell.fill = PatternFill("solid", fgColor=PRIORITY_COLORS[n.priority])
        pcell.alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=9).number_format = "yyyy-mm-dd"
        ws.row_dimensions[r].height = 48
    if needs:
        status = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
        ws.add_data_validation(status)
        status.add(f"J8:J{r}")
    else:
        ws.cell(row=8, column=1, value="No issues found in the numbers provided.")

    ws.cell(row=r + 2, column=1, value=("Recommendations are generated from the financial ratios using rules of "
                                        "thumb. Review them with the client before acting on them."))
    ws.cell(row=r + 2, column=1).font = theme.subtitle_font
    for col, width in zip("ABCDEFGHIJ", (5, 11, 18, 36, 44, 30, 24, 14, 12, 13)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A8"
    return ws


def _stamp_license(wb, ws, note, theme):
    """Put the buyer's license line under the sheet content and in the file properties."""
    if not note:
        return
    cell = ws.cell(row=ws.max_row + 2, column=1, value=note)
    cell.font = theme.subtitle_font
    wb.properties.description = note
    wb.properties.keywords = note


def build_company_workbook(analysis, path, brand_color="1F3A5F", prepared_by=None, pro=False,
                           license_note=None):
    theme = Theme(brand_color)
    subtitle = _subtitle(analysis, prepared_by)
    wb = Workbook()
    wb.active.title = "Dashboard"  # placeholder; filled last so charts can reference other sheets
    fin_ws, fin_rows = _financials_sheet(wb, analysis, theme, subtitle)
    ratio_ws, ratio_rows = _ratios_sheet(wb, analysis, fin_rows, theme, subtitle)
    _forecast_sheet(wb, analysis, fin_rows, theme, subtitle)
    n = len(analysis.company.years)
    charts = [_performance_chart(fin_ws, fin_rows, n), _margin_chart(ratio_ws, ratio_rows, n)]
    _dashboard_sheet(wb, analysis, theme, subtitle, charts)
    if pro:
        _client_needs_sheet(wb, analysis, theme, subtitle)
    wb.properties.title = f"{analysis.company.name} — Company Analysis"
    if prepared_by and safe_text(prepared_by):
        wb.properties.creator = safe_text(prepared_by)
    _stamp_license(wb, wb["Dashboard"], license_note, theme)
    wb.save(path)
    return path


def build_comparison_workbook(analyses, path, brand_color="1F3A5F", prepared_by=None, license_note=None):
    theme = Theme(brand_color)
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    _title(ws, "Peer Comparison", _subtitle(None, prepared_by), theme)
    headers = ["Rank", "Company", "Latest Year", "Revenue", "Revenue CAGR", "Gross Margin",
               "Net Margin", "ROE", "Current Ratio", "Debt / Equity", "FCF Margin", "Score", "Grade"]
    _header_row(ws, 4, headers, theme)
    ranked = sorted(analyses, key=lambda a: a.score if a.score is not None else -1, reverse=True)
    for i, a in enumerate(ranked, start=1):
        r = 4 + i
        latest = a.ratios[-1]
        values = [i, a.company.name, a.company.years[-1], a.company.rows[-1]["revenue"], a.revenue_cagr,
                  latest.get("gross_margin"), latest.get("net_margin"), latest.get("roe"),
                  latest.get("current_ratio"), latest.get("debt_to_equity"), latest.get("fcf_margin"),
                  a.score, a.grade]
        fmts = [None, None, "0", MONEY, PCT, PCT, PCT, PCT, MULT, MULT, PCT, "0.0", None]
        for c, (v, fmt) in enumerate(zip(values, fmts), start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = theme.border
            if fmt:
                cell.number_format = fmt
            if r % 2 == 0:
                cell.fill = theme.band_fill
        grade_cell = ws.cell(row=r, column=13)
        grade_cell.font = Font(bold=True, color="FFFFFF")
        grade_cell.fill = PatternFill("solid", fgColor=GRADE_COLORS[a.grade])
        grade_cell.alignment = Alignment(horizontal="center")
    widths = [7, 24, 11, 16, 13, 13, 12, 10, 13, 13, 12, 9, 8]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C5"

    last = 4 + len(ranked)
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Financial Health Score"
    chart.height, chart.width = 8, 16
    chart.add_data(Reference(ws, min_col=12, min_row=4, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=5, max_row=last))
    chart.legend = None
    ws.add_chart(chart, f"B{last + 3}")
    _stamp_license(wb, ws, license_note, theme)
    wb.save(path)
    return path


# ======================================================================
# WEB
# ======================================================================
SAMPLE_DATA = {
    'brightpath_software.csv': 'year,revenue,cogs,operating_expenses,interest_expense,taxes,net_income,total_assets,total_liabilities,shareholder_equity,current_assets,current_liabilities,cash,operating_cash_flow,capex\n2021,2.1M,420K,2.4M,0,0,-720K,5.5M,1.8M,3.7M,4.6M,1.1M,4.1M,-310K,90K\n2022,3.9M,740K,3.6M,0,0,-440K,7.2M,2.4M,4.8M,6.0M,1.5M,5.2M,120K,140K\n2023,6.4M,1.15M,4.6M,0,90K,560K,9.6M,3.1M,6.5M,8.1M,2.0M,6.9M,1.4M,210K\n2024,9.8M,1.67M,6.1M,0,410K,1.62M,13.4M,4.0M,9.4M,11.5M,2.6M,9.6M,2.9M,320K\n',
    'northwind_coffee.csv': 'year,revenue,cogs,operating_expenses,interest_expense,taxes,net_income,total_assets,total_liabilities,shareholder_equity,current_assets,current_liabilities,cash,operating_cash_flow,capex\n2020,4200000,1890000,1470000,42000,180000,618000,3100000,1200000,1900000,1250000,620000,540000,710000,240000\n2021,4950000,2178000,1683000,40000,232000,817000,3600000,1250000,2350000,1480000,650000,690000,905000,260000\n2022,5800000,2552000,1972000,38000,280000,958000,4200000,1300000,2900000,1720000,700000,860000,1080000,310000\n2023,6650000,2860000,2261000,35000,330000,1164000,4900000,1350000,3550000,2010000,740000,1050000,1290000,340000\n2024,7480000,3141600,2543200,32000,380000,1383200,5650000,1400000,4250000,2320000,780000,1270000,1520000,370000\n',
    'summit_logistics.csv': 'Fiscal Year,Sales,Cost of Sales,OpEx,Interest,Income Tax,Net Profit,Total Assets,Total Liabilities,Equity,Current Assets,Current Liabilities,Cash,Cash From Operations,CapEx\n2020,"$18,500,000","$14,060,000","$3,330,000","$610,000","$120,000","$380,000","$22,000,000","$15,400,000","$6,600,000","$5,200,000","$4,900,000","$900,000","$1,450,000","$1,900,000"\n2021,"$19,200,000","$14,784,000","$3,456,000","$650,000","$90,000","$220,000","$23,100,000","$16,500,000","$6,600,000","$5,100,000","$5,300,000","$700,000","$1,200,000","$2,100,000"\n2022,"$19,900,000","$15,522,000","$3,582,000","$720,000","$20,000","$56,000","$24,400,000","$17,800,000","$6,600,000","$5,000,000","$5,700,000","$520,000","$980,000","$2,300,000"\n2023,"$19,300,000","$15,247,000","$3,474,000","$780,000","$0","($201,000)","$24,900,000","$18,500,000","$6,400,000","$4,800,000","$6,000,000","$410,000","$760,000","$1,800,000"\n2024,"$18,700,000","$14,773,000","$3,366,000","$820,000","$0","($259,000)","$25,200,000","$19,100,000","$6,100,000","$4,600,000","$6,200,000","$350,000","$690,000","$1,500,000"\n',
}


def _slug(name):
    return "".join(ch if ch.isalnum() else "_" for ch in name).strip("_").lower() or "company"


def _summary(a, path):
    return {
        "name": a.company.name,
        "years": [a.company.years[0], a.company.years[-1]],
        "score": a.score,
        "grade": a.grade,
        "insights": a.insights,
        "warnings": a.warnings,
        "file": path.name,
    }


def run(request_json):
    """Build reports from a JSON request and return a JSON reply.

    Request: {"groups": [{"files": [paths], "name": str|None}], "out_dir": str, "pro": bool,
              "brand_color": str, "prepared_by": str|None, "grades": str|None,
              "license_note": str|None}
    Each group is one company: one file, or several accounting exports to combine.
    Reply: {"ok": true, "companies": [...], "files": [...], "zip": name} or {"ok": false, "error": str}
    """
    req = json.loads(request_json)
    out_dir = Path(req["out_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)
    pro = bool(req.get("pro"))
    try:
        cutoffs = parse_grades(req["grades"]) if pro and req.get("grades") else DEFAULT_GRADES
        analyses, companies, files = [], [], []
        for group in req["groups"]:
            paths = [Path(p) for p in group["files"]]
            name = (group.get("name") or "").strip() or None
            if len(paths) == 1:
                company = load_file(paths[0], name)
            else:
                company = load_company(paths, name, default_name="Company")
            a = analyze(company, grade_cutoffs=cutoffs)
            path = out_dir / f"{_slug(company.name)}_analysis.xlsx"
            build_company_workbook(a, path, req.get("brand_color") or "1F3A5F", req.get("prepared_by") or None,
                                   pro=pro, license_note=req.get("license_note"))
            analyses.append(a)
            companies.append(_summary(a, path))
            files.append(path.name)
        if len(analyses) > 1:
            path = out_dir / "peer_comparison.xlsx"
            build_comparison_workbook(analyses, path, req.get("brand_color") or "1F3A5F",
                                      req.get("prepared_by") or None, license_note=req.get("license_note"))
            files.append(path.name)
        zip_name = None
        if len(files) > 1:
            zip_name = "flash_analyst_reports.zip"
            with zipfile.ZipFile(out_dir / zip_name, "w", zipfile.ZIP_DEFLATED) as zf:
                for name in files:
                    zf.write(out_dir / name, name)
        return json.dumps({"ok": True, "companies": companies, "files": files, "zip": zip_name})
    except (DataError, ValueError) as exc:
        return json.dumps({"ok": False, "error": str(exc)})


def sample_files(out_dir):
    """Write the bundled sample companies to out_dir and return their paths, for the demo."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for filename, content in SAMPLE_DATA.items():
        dest = out_dir / filename
        dest.write_text(content)
        paths.append(str(dest))
    return json.dumps(paths)
